"""文件锻造纯逻辑回归（零 LLM 零 DB）——格式合法性是造数中心的命门。

守护：八格式各自产出**格式合法**的字节（魔数/可被对应库打开）；kind 归一兜底；
大文件定容（txt 精确、pdf 逼近）；超限抛错。
"""

import io

import pytest

from app.services.test_data.file_forge import (
    MAX_FORGE_BYTES,
    forge_file,
    mime_of,
    normalize_kind,
)


class TestNormalizeKind:
    def test_explicit_kind(self):
        assert normalize_kind("PDF") == "pdf"
        assert normalize_kind(".xlsx") == "xlsx"

    def test_fallback_to_ext(self):
        assert normalize_kind("", "a.docx") == "docx"
        assert normalize_kind("bogus", "a.csv") == "csv"

    def test_default_txt(self):
        assert normalize_kind("", "noext") == "txt"
        assert normalize_kind("unknown", "x.zzz") == "txt"


class TestForgeFormats:
    def test_pdf_valid(self):
        import fitz

        content, mime = forge_file({"kind": "pdf", "file_name": "a.pdf", "text": "合同\n金额 100"})
        assert mime == "application/pdf"
        d = fitz.open(stream=content, filetype="pdf")
        assert d.page_count >= 1

    def test_docx_valid(self):
        from docx import Document

        content, _ = forge_file({"kind": "docx", "file_name": "a.docx", "text": "l1\nl2\nl3"})
        assert len(Document(io.BytesIO(content)).paragraphs) == 3

    def test_xlsx_valid(self):
        from openpyxl import load_workbook

        content, _ = forge_file(
            {"kind": "xlsx", "file_name": "a.xlsx", "table": {"headers": ["a", "b"], "rows": [[1, 2], [3, 4]]}}
        )
        wb = load_workbook(io.BytesIO(content))
        assert wb.active.max_row == 3  # 表头 + 2 行

    def test_csv_bom(self):
        content, _ = forge_file({"kind": "csv", "file_name": "a.csv", "table": {"headers": ["x"], "rows": [["y"]]}})
        assert content[:3] == b"\xef\xbb\xbf"  # utf-8-sig，Excel 中文不乱码

    def test_png_magic(self):
        content, mime = forge_file({"kind": "png", "file_name": "a.png", "text": "hi"})
        assert mime == "image/png"
        assert content[:8] == b"\x89PNG\r\n\x1a\n"

    def test_json_valid(self):
        import json

        content, _ = forge_file({"kind": "json", "file_name": "a.json", "text": '{"k":"v"}'})
        assert json.loads(content) == {"k": "v"}

    def test_json_wraps_non_json(self):
        import json

        content, _ = forge_file({"kind": "json", "file_name": "a.json", "text": "plain"})
        assert json.loads(content) == {"content": "plain"}

    def test_txt_plain(self):
        content, mime = forge_file({"kind": "txt", "file_name": "a.txt", "text": "hello"})
        assert content == b"hello" and mime == "text/plain"


class TestSizeControl:
    def test_txt_exact_size(self):
        content, _ = forge_file({"kind": "txt", "file_name": "a.txt", "text": "x", "target_size_bytes": 50000})
        assert len(content) == 50000

    def test_pdf_approaches_target(self):
        content, _ = forge_file({"kind": "pdf", "file_name": "a.pdf", "text": "big", "target_size_bytes": 1_500_000})
        assert len(content) >= 1_500_000

    def test_over_limit_raises(self):
        with pytest.raises(ValueError):
            forge_file({"kind": "bin", "file_name": "a.bin", "target_size_bytes": MAX_FORGE_BYTES + 1})


def test_mime_of():
    assert mime_of("pdf") == "application/pdf"
    assert mime_of("unknown") == "application/octet-stream"
