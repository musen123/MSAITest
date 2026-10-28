"""文档媒体无损提取共享层回归（doc_media，设计见 docs/文档解析-图片与表格无损化-设计方案.md）

守护的硬约束：
  - **无图文档输出与旧实现逐字节等价**（docx 双渲染 / pdf 渲染，legacy 算法
    内联为对照组）——存量行为零变化是共享层替换的前提；
  - 提取顺序 = 文档顺序（段落→段内图→表格 / 页文本→页内图）；
  - 六道护栏：小图过滤、张数上限、md5 去重、单图失败占位、无 VISION 降级、
    描述限长单行化。
"""

from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest

from app.services.common.doc_media import (
    MAX_DESC_CHARS,
    describe_images,
    extract_docx_segments,
    extract_pdf_segments,
    render_markdown_docx,
    render_pdf,
    render_plain,
)

# ─────────────────────────────────────────────────────────────
# fixture 生成（运行时现做，不落二进制文件进仓库）
# ─────────────────────────────────────────────────────────────


def _png_bytes(width: int, height: int, seed: int = 7) -> bytes:
    """随机噪声 PNG（不可压缩 → 大图稳超 5KB 字节阈值；seed 变化产生不同 md5）。"""
    import random

    rnd = random.Random(seed)
    samples = bytes(rnd.getrandbits(8) for _ in range(width * height * 3))
    pix = fitz.Pixmap(fitz.csRGB, width, height, samples, False)
    return pix.tobytes("png")


@pytest.fixture()
def docx_with_media(tmp_path: Path) -> dict:
    """docx：标题 + 段落 + 大图×2(同图) + 小图 + 表格（顺序敏感）"""
    import docx

    doc = docx.Document()
    doc.add_heading("需求规格", level=1)
    doc.add_paragraph("登录模块的说明文字。")
    big = _png_bytes(400, 300)
    small = _png_bytes(32, 32)
    big_path, small_path = tmp_path / "big.png", tmp_path / "small.png"
    big_path.write_bytes(big)
    small_path.write_bytes(small)
    doc.add_picture(str(big_path))
    doc.add_paragraph("下面是校验规则表。")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "字段"
    table.rows[0].cells[1].text = "规则"
    table.rows[1].cells[0].text = "手机号"
    table.rows[1].cells[1].text = "11 位"
    doc.add_paragraph("重复插入同一张图：")
    doc.add_picture(str(big_path))  # 同图第二次 → md5 去重
    doc.add_picture(str(small_path))  # 小图 → 过滤
    path = tmp_path / "fixture.docx"
    doc.save(str(path))
    return {"path": path, "big_bytes": len(big)}


@pytest.fixture()
def docx_text_only(tmp_path: Path) -> Path:
    import docx

    doc = docx.Document()
    doc.add_heading("标题一", level=1)
    doc.add_paragraph("第一段。")
    doc.add_heading("标题二", level=2)
    doc.add_paragraph("第二段。")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "a"
    table.rows[0].cells[1].text = "b"
    path = tmp_path / "plain.docx"
    doc.save(str(path))
    return path


@pytest.fixture()
def pdf_with_image(tmp_path: Path) -> Path:
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "First page text.")
    img = _png_bytes(300, 200)
    page1.insert_image(fitz.Rect(72, 100, 372, 300), stream=img)
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Second page text.")
    path = tmp_path / "fixture.pdf"
    doc.save(str(path))
    doc.close()
    return path


class FakeVision:
    def __init__(self, text: str = "图中为登录页,含手机号输入框(11位)与密码框。"):
        self.text = text
        self.calls = 0

    async def ainvoke(self, messages):
        self.calls += 1
        return SimpleNamespace(content=self.text)


class BoomVision:
    async def ainvoke(self, messages):
        raise RuntimeError("vlm down")


# ─────────────────────────────────────────────────────────────
# 等价性（legacy 算法内联对照）
# ─────────────────────────────────────────────────────────────


def _legacy_read_docx_plain(path: Path) -> str:
    """旧 func_test._read_docx：只读段落，\\n join 非空段。"""
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _legacy_docx_to_markdown(path: Path) -> str:
    """旧 kb_parse_service._docx_to_markdown 逐字复刻（对照组）。"""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    lines: list[str] = []
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name or "") if item.style is not None else ""
            if style.startswith(("Heading", "标题")):
                digits = "".join(ch for ch in style if ch.isdigit())
                level = min(max(int(digits), 1), 6) if digits else 1
                lines.append(f"{'#' * level} {text}")
            else:
                lines.append(text)
        elif isinstance(item, Table):
            rows = [
                [c.text.strip().replace("\n", " ").replace("|", "\\|") for c in row.cells]
                for row in item.rows
            ]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("|" + "---|" * width)
            for r in rows[1:]:
                lines.append("| " + " | ".join(r) + " |")
            lines.append("")
    return "\n".join(lines).strip()


class TestByteEquivalence:
    async def test_docx_markdown_render_equals_legacy(self, docx_text_only) -> None:
        segments = extract_docx_segments(docx_text_only)
        segments, _ = await describe_images(segments, None)
        assert render_markdown_docx(segments) == _legacy_docx_to_markdown(
            docx_text_only
        )

    async def test_docx_plain_render_equals_legacy_for_paragraphs(
        self, docx_text_only
    ) -> None:
        # A 链旧输出只有段落；新输出=段落+表格。剔除表格段后须逐字节等价
        segments = extract_docx_segments(docx_text_only)
        para_only = [s for s in segments if s.kind == "para"]
        assert render_plain(para_only) == _legacy_read_docx_plain(docx_text_only)

    async def test_pdf_render_equals_legacy(self, pdf_with_image) -> None:
        # 无 VISION → 图段被移除，渲染须与旧 _read_pdf(逐页 get_text \n join) 等价
        doc = fitz.open(str(pdf_with_image))
        legacy = "\n".join(page.get_text() for page in doc)
        doc.close()
        segments = extract_pdf_segments(pdf_with_image)
        segments, _ = await describe_images(segments, None)
        assert render_pdf(segments) == legacy


# ─────────────────────────────────────────────────────────────
# 提取顺序与内容
# ─────────────────────────────────────────────────────────────


class TestExtraction:
    def test_docx_order_and_kinds(self, docx_with_media) -> None:
        segments = extract_docx_segments(docx_with_media["path"])
        kinds = [s.kind for s in segments]
        # 三张图(大图×2+小图)全部被提取(过滤发生在描述阶段)
        assert kinds.count("image") == 3
        assert kinds.count("table") == 1
        # 顺序：表格出现在第一张图之后、第二张图之前
        first_img = kinds.index("image")
        table_at = kinds.index("table")
        assert first_img < table_at < len(kinds) - 1

    def test_pdf_pages_and_images(self, pdf_with_image) -> None:
        segments = extract_pdf_segments(pdf_with_image)
        pages = [s for s in segments if s.kind == "page_text"]
        images = [s for s in segments if s.kind == "image"]
        assert len(pages) == 2 and len(images) == 1
        assert images[0].page == 1
        # 图段紧跟第 1 页文本之后
        assert segments.index(images[0]) == segments.index(pages[0]) + 1


# ─────────────────────────────────────────────────────────────
# 护栏
# ─────────────────────────────────────────────────────────────


class TestGuardrails:
    async def test_dedupe_and_small_filter(self, docx_with_media) -> None:
        fake = FakeVision()
        segments = extract_docx_segments(docx_with_media["path"])
        segments, stats = await describe_images(segments, fake)
        # 大图×2 同 md5 → 只调一次 VLM;小图被过滤
        assert fake.calls == 1
        assert stats.described == 1 and stats.deduped == 1
        assert stats.skipped_small == 1
        rendered = render_markdown_docx(segments)
        assert rendered.count("> 【图片内容】") == 2  # 两处大图均回填(复用描述)
        assert "手机号输入框" in rendered

    async def test_no_vision_degrades(self, docx_with_media) -> None:
        segments = extract_docx_segments(docx_with_media["path"])
        segments, stats = await describe_images(segments, None)
        assert stats.skipped_no_vision == 3
        assert all(s.kind != "image" for s in segments)

    async def test_failure_placeholder(self, docx_with_media) -> None:
        segments = extract_docx_segments(docx_with_media["path"])
        segments, stats = await describe_images(segments, BoomVision())
        assert stats.failed >= 1
        assert any("图片解析失败" in s.text for s in segments if s.kind == "image")

    async def test_cap(self, pdf_with_image, tmp_path) -> None:
        # 造 3 页各 1 张不同图的 pdf,cap=2 → 第 3 张跳过并留痕
        doc = fitz.open()
        for i in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), f"page {i}")
            page.insert_image(
                fitz.Rect(72, 100, 372, 300),
                stream=_png_bytes(300, 200, seed=100 + i),
            )
        p = tmp_path / "cap.pdf"
        doc.save(str(p))
        doc.close()
        segments = extract_pdf_segments(p)
        fake = FakeVision()
        segments, stats = await describe_images(segments, fake, max_images=2)
        assert fake.calls == 2
        assert stats.skipped_over_cap == 1
        assert any("未解析" in n for n in stats.notes)

    async def test_desc_single_line_and_capped(self, docx_with_media) -> None:
        long_text = "第一行\n第二行 " * 200
        segments = extract_docx_segments(docx_with_media["path"])
        segments, _ = await describe_images(segments, FakeVision(long_text))
        img = next(s for s in segments if s.kind == "image" and s.text)
        assert "\n" not in img.text and len(img.text) <= MAX_DESC_CHARS

    async def test_page_tag_in_pdf_block(self, pdf_with_image) -> None:
        segments = extract_pdf_segments(pdf_with_image)
        segments, _ = await describe_images(segments, FakeVision())
        rendered = render_pdf(segments)
        assert "> 【图片内容|第1页】" in rendered


# ─────────────────────────────────────────────────────────────
# P2 扫描件转写
# ─────────────────────────────────────────────────────────────


@pytest.fixture()
def scanned_pdf(tmp_path: Path) -> Path:
    """无文本层 PDF：每页只有一张整页图（模拟扫描件）。"""
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page(width=400, height=300)
        page.insert_image(
            fitz.Rect(0, 0, 400, 300), stream=_png_bytes(200, 150, seed=200 + i)
        )
    path = tmp_path / "scanned.pdf"
    doc.save(str(path))
    doc.close()
    return path


class TestScannedPdf:
    async def test_transcribe_pages(self, scanned_pdf) -> None:
        from app.services.common.doc_media import transcribe_scanned_pdf

        fake = FakeVision("## 第x页标题\n正文内容")
        md, pages, over_cap = await transcribe_scanned_pdf(scanned_pdf, fake)
        assert pages == 3 and not over_cap and fake.calls == 3
        assert md.count("正文内容") == 3

    async def test_transcribe_page_cap(self, scanned_pdf) -> None:
        from app.services.common.doc_media import transcribe_scanned_pdf

        fake = FakeVision()
        md, pages, over_cap = await transcribe_scanned_pdf(
            scanned_pdf, fake, max_pages=2
        )
        assert pages == 2 and over_cap and fake.calls == 2

    async def test_transcribe_page_failure_placeholder(self, scanned_pdf) -> None:
        from app.services.common.doc_media import transcribe_scanned_pdf

        md, pages, _ = await transcribe_scanned_pdf(scanned_pdf, BoomVision())
        assert pages == 3
        assert md.count("转写失败") == 3

    async def test_transcribe_requires_vision(self, scanned_pdf) -> None:
        from app.services.common.doc_media import transcribe_scanned_pdf

        with pytest.raises(ValueError, match="VISION"):
            await transcribe_scanned_pdf(scanned_pdf, None)

    async def test_parse_file_scanned_switch(self, scanned_pdf, monkeypatch) -> None:
        """kb 解析入口：无文本层 PDF 关开关报可读错误、开开关走转写。"""
        from app.services.knowledge import kb_parse_service as kps

        monkeypatch.setattr(
            kps, "_load_to_local", lambda _k: (scanned_pdf, None)
        )
        with pytest.raises(ValueError, match="扫描"):
            await kps.parse_file_for_kb(
                storage_key="x", file_ext=".pdf", file_name="s.pdf",
                text_llm=object(), vision_llm=FakeVision(),
            )
        res = await kps.parse_file_for_kb(
            storage_key="x", file_ext=".pdf", file_name="s.pdf",
            text_llm=object(), vision_llm=FakeVision("第页内容"),
            scanned_pdf=True,
        )
        assert res.parse_mode == "llm_normalized"
        assert "扫描件已逐页转写 3 页" in (res.media_note or "")
